# -*- coding: utf-8 -*-
"""Тесты записи Excel (с колонкой уверенности) и формирования лог-файла."""

from pathlib import Path

from openpyxl import load_workbook

from tn_parser import (
    FieldConfidence,
    ParsedRow,
    build_log_lines,
    write_excel,
    write_log,
)
from tn_parser.excel import COLUMNS


def _sample_rows():
    high = ParsedRow(
        waybill="Транспортная накладная № 001",
        date="01.01.2024", number="001",
        shipper="ООО Ромашка", consignee="ООО Василёк",
        cargo="Мука", carrier="ООО Быстрые",
        vehicle="А 123 ВС 777", reception="г. Москва 01.01.2024",
        source="a.pdf", note="",
        confidence=FieldConfidence(
            date=1.0, number=0.9, shipper=0.9, consignee=0.9,
            cargo=0.9, carrier=0.9, vehicle=1.0, reception=0.9,
        ),
    )
    low = ParsedRow(
        waybill="Транспортная накладная",
        date="отсутствует", number="отсутствует",
        shipper="отсутствует", consignee="отсутствует",
        cargo="отсутствует", carrier="отсутствует",
        vehicle="отсутствует", reception="отсутствует",
        source="b.pdf", note="LOW_CONF",
        confidence=FieldConfidence(),
    )
    err = ParsedRow.empty_missing("c.pdf", note="ERROR: boom")
    return high, low, err


class TestExcel:
    def test_has_12_columns_including_confidence(self, tmp_path):
        rows = list(_sample_rows())
        out = tmp_path / "out.xlsx"
        write_excel(rows, str(out))

        wb = load_workbook(out)
        ws = wb.active
        headers = [ws.cell(row=1, column=i + 1).value for i in range(len(COLUMNS))]
        assert len(headers) == 12
        assert headers[-1] == "Уверенность, %"

    def test_confidence_values(self, tmp_path):
        high, low, err = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high, low, err], str(out))

        wb = load_workbook(out)
        ws = wb.active
        # Колонка confidence — 12-я.
        assert ws.cell(row=2, column=12).value == 93  # high ≈ 0.925 → 93%
        assert ws.cell(row=3, column=12).value == 0   # low
        assert ws.cell(row=4, column=12).value == 0   # err

    def test_confidence_fill_colors(self, tmp_path):
        high, low, _ = _sample_rows()
        out = tmp_path / "out.xlsx"
        write_excel([high, low], str(out))
        wb = load_workbook(out)
        ws = wb.active
        # Высокая уверенность — зелёная; нулевая — красная.
        high_fill = ws.cell(row=2, column=12).fill.fgColor.rgb
        low_fill = ws.cell(row=3, column=12).fill.fgColor.rgb
        assert "D9EAD3" in (high_fill or "").upper()
        assert "F4CCCC" in (low_fill or "").upper()


class TestReport:
    def test_build_log_lines_structure(self):
        high, low, err = _sample_rows()
        lines = build_log_lines(
            input_path="/tmp/pdfs",
            output_path="/tmp/out.xlsx",
            elapsed_s=1.23,
            rows_by_file={"a.pdf": [high], "b.pdf": [low], "c.pdf": [err]},
        )
        text = "\n".join(lines)
        assert "Парсер транспортных накладных" in text
        assert "/tmp/pdfs" in text
        assert "[OK  ]" in text
        assert "[WARN]" in text
        assert "[ERR ]" in text
        assert "93%" in text
        # Для low confidence должны быть перечислены проблемные поля.
        assert "Проверить:" in text
        # Для ошибки должен быть виден её текст.
        assert "ERROR: boom" in text

    def test_write_log_creates_file(self, tmp_path):
        lines = ["строка один", "строка два"]
        target = tmp_path / "extraction.log"
        write_log(str(target), lines)
        content = target.read_text(encoding="utf-8")
        assert "строка один\nстрока два\n" == content
